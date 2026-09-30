from io import BytesIO
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from reportlab.lib.styles import getSampleStyleSheet

def resume_docx(x):
    d=Document(); p=x.resume.personal; d.add_heading(p.name or "Resume",0)
    contact=" | ".join(v for v in [p.email,p.phone,p.location,p.linkedin,p.github] if v)
    if contact:d.add_paragraph(contact)
    d.add_heading("Professional Summary",1); d.add_paragraph(x.optimized_resume.summary)
    d.add_heading("Experience",1)
    for e in x.resume.experience:
        d.add_paragraph(f"{e.role} — {e.company}",style="Heading 2")
        for b in e.bullets:d.add_paragraph(b,style="List Bullet")
    d.add_heading("Projects",1)
    for z in x.resume.projects:
        d.add_paragraph(z.name,style="Heading 2")
        for b in z.bullets or [z.description]:d.add_paragraph(b,style="List Bullet")
    d.add_heading("Skills",1);d.add_paragraph(", ".join(x.resume.skills.technical+x.resume.skills.tools+x.resume.skills.soft))
    d.add_heading("Education",1)
    for e in x.resume.education:d.add_paragraph(f"{e.degree} {e.field} — {e.institution} ({e.start_date}–{e.end_date})")
    o=BytesIO();d.save(o);return o.getvalue()

def cover_docx(x):
    d=Document();d.add_heading(x.cover_letter.subject or "Cover Letter",0);d.add_paragraph(x.cover_letter.body)
    o=BytesIO();d.save(o);return o.getvalue()

def report_pdf(x):
    o=BytesIO();d=SimpleDocTemplate(o,pagesize=A4);s=getSampleStyleSheet()
    d.build([Paragraph("CareerCraft AI — Analysis Report",s["Title"]),Spacer(1,12),
             Paragraph(f"Alignment: {x.match.alignment_score}/100",s["Heading2"]),
             Paragraph(x.match.explanation,s["BodyText"]),Spacer(1,10),
             Paragraph("Missing skills",s["Heading2"]),Paragraph(", ".join(x.skill_gap.missing) or "None",s["BodyText"])])
    return o.getvalue()
